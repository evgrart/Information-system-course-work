"""Build the stage-one report and editable UML diagrams without external services."""
from pathlib import Path
import re
import math
import html
from datetime import datetime, timezone
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
                               Image as PDFImage, PageBreak, KeepTogether, Preformatted)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / 'docs' / 'part1'
DIAGRAMS = REPORT_DIR / 'uml'
DIAGRAMS.mkdir(parents=True,exist_ok=True)
FONTS = Path('C:/Windows/Fonts')
INK = '#263746'
BLUE = '#edf3f8'


class Drawing:
    def __init__(self, name, width, height, title):
        self.name, self.width, self.height = name, width, height
        self.im = Image.new('RGB', (width, height), 'white')
        self.d = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                    '<rect width="100%" height="100%" fill="white"/>']
        self.text(width / 2, 35, title, 29, bold=True)

    def line(self, points, dashed=False, arrow=False, triangle=False, color=INK):
        for a, b in zip(points, points[1:]):
            if dashed:
                dist = math.hypot(b[0]-a[0], b[1]-a[1])
                for start in range(0, int(dist), 16):
                    end = min(start+9, dist)
                    self.d.line([(a[0]+(b[0]-a[0])*start/dist, a[1]+(b[1]-a[1])*start/dist),
                                 (a[0]+(b[0]-a[0])*end/dist, a[1]+(b[1]-a[1])*end/dist)], fill=color, width=3)
            else:
                self.d.line([a,b], fill=color, width=3)
        pts=' '.join(f'{x},{y}' for x,y in points)
        self.svg.append(f'<polyline points="{pts}" fill="none" stroke="{color}" stroke-width="3"'+(' stroke-dasharray="9 7"' if dashed else '')+'/>')
        if arrow or triangle:
            a,b=points[-2:]
            angle=math.atan2(b[1]-a[1],b[0]-a[0])
            ends=[(b[0]-21*math.cos(angle-s),b[1]-21*math.sin(angle-s)) for s in (-.45,.45)]
            if triangle:
                self.d.polygon([b,*ends],fill='white',outline=color,width=3)
                self.svg.append(f'<polygon points="{b[0]},{b[1]} {ends[0][0]},{ends[0][1]} {ends[1][0]},{ends[1][1]}" fill="white" stroke="{color}" stroke-width="3"/>')
            else:
                self.d.line([ends[0],b,ends[1]],fill=color,width=3)
                self.svg.append(f'<polyline points="{ends[0][0]},{ends[0][1]} {b[0]},{b[1]} {ends[1][0]},{ends[1][1]}" fill="none" stroke="{color}" stroke-width="3"/>')

    def text(self,x,y,value,size=26,bold=False):
        font=ImageFont.truetype(str(FONTS / ('arialbd.ttf' if bold else 'arial.ttf')),size)
        lines=value.split('\n')
        for i,line in enumerate(lines):
            yy=y+(i-(len(lines)-1)/2)*(size+7)
            self.d.text((x,yy),line,font=font,fill=INK,anchor='mm')
            self.svg.append(f'<text x="{x}" y="{yy}" font-family="Arial, sans-serif" font-size="{size}" font-weight="'+('bold' if bold else 'normal')+'" text-anchor="middle" dominant-baseline="central" fill="'+INK+'">'+html.escape(line)+'</text>')

    def box(self,x1,y1,x2,y2,label='',fill='white',size=26):
        self.d.rectangle((x1,y1,x2,y2),fill=fill,outline=INK,width=3)
        self.svg.append(f'<rect x="{x1}" y="{y1}" width="{x2-x1}" height="{y2-y1}" fill="{fill}" stroke="{INK}" stroke-width="3"/>')
        if label:self.text((x1+x2)/2,(y1+y2)/2,label,size)

    def uc(self,x,y,label,w=430,h=115):
        self.d.ellipse((x-w/2,y-h/2,x+w/2,y+h/2),fill=BLUE,outline=INK,width=3)
        self.svg.append(f'<ellipse cx="{x}" cy="{y}" rx="{w/2}" ry="{h/2}" fill="{BLUE}" stroke="{INK}" stroke-width="3"/>')
        self.text(x,y,label,32)

    def actor(self,x,y,label):
        self.d.ellipse((x-20,y-60,x+20,y-20),outline=INK,width=3)
        self.svg.append(f'<circle cx="{x}" cy="{y-40}" r="20" fill="white" stroke="{INK}" stroke-width="3"/>')
        self.line([(x,y-20),(x,y+35)])
        self.line([(x-35,y+3),(x+35,y+3)])
        self.line([(x,y+35),(x-30,y+75)])
        self.line([(x,y+35),(x+30,y+75)])
        self.text(x,y+108,label,28)

    def save(self):
        self.im.save(DIAGRAMS / (self.name+'.png'))
        (DIAGRAMS / (self.name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf-8')


def uml_file(name,body):
    (DIAGRAMS/(name+'.puml')).write_text('@startuml\nleft to right direction\nskinparam shadowing false\nskinparam defaultFontName Arial\n'+body+'\n@enduml\n',encoding='utf-8')


def make_diagrams():
    g=Drawing('01_access',1650,1180,'ИС «Потеряшки»: регистрация и подписка')
    g.box(340,85,1290,1100)
    g.text(815,110,'Граница информационной системы',22)
    g.actor(135,200,'Посетитель')
    g.actor(135,440,'Владелец\nаккаунта')
    g.actor(135,970,'Подписчик')
    g.actor(1490,200,'Почтовый\nсервис')
    g.actor(1490,440,'Модератор')
    g.actor(1490,920,'Платёжный\nсервис')
    g.line([(170,203),(290,203),(290,200),(365,200)])
    for y in (440,680,920):g.line([(170,443),(290,443),(290,y),(365,y)])
    g.line([(135,910),(135,590)],triangle=True)
    g.uc(580,200,'UC-01. Регистрация\nи подтверждение почты')
    g.uc(580,440,'UC-02. Проверка\nпрофиля')
    g.uc(580,680,'UC-03. Вход')
    g.uc(580,920,'UC-04. Подписка')
    g.uc(1075,680,'Восстановить\nпароль',w=340)
    g.uc(1075,920,'Обработать\nрезультат оплаты',w=340)
    g.line([(795,200),(1455,200)])
    g.line([(795,440),(1455,440)])
    g.line([(1245,680),(1340,680),(1340,203),(1455,203)])
    g.line([(1245,920),(1455,920)])
    g.line([(905,680),(795,680)],dashed=True,arrow=True)
    g.text(845,625,'«extend»',21)
    g.text(1075,778,'[пароль утрачен]',21)
    g.line([(795,920),(905,920)],dashed=True,arrow=True)
    g.text(845,865,'«include»',21)
    g.save()
    uml_file('01_access','''actor "Посетитель" as Visitor
actor "Владелец аккаунта" as Account
actor "Подписчик" as Subscriber
actor "Модератор" as Moderator
actor "Почтовый сервис" as Mail
actor "Платёжный сервис" as Payment
Subscriber --|> Account
rectangle "ИС Потеряшки" {
usecase "UC-01. Регистрация и подтверждение почты" as U1
usecase "UC-02. Проверка профиля" as U2
usecase "UC-03. Вход" as U3
usecase "UC-04. Подписка" as U4
usecase "Восстановить пароль" as Reset
usecase "Обработать результат оплаты" as Result
}
Visitor -- U1
Account -- U2
Account -- U3
Account -- U4
Mail -- U1
Mail -- Reset
Moderator -- U2
Payment -- Result
Reset ..> U3 : <<extend>>\n[пароль утрачен]
U4 ..> Result : <<include>>''')

    g=Drawing('02_return',1720,1280,'ИС «Потеряшки»: поиск и возврат')
    g.box(345,85,1350,1190)
    g.text(845,110,'Граница информационной системы',22)
    g.actor(140,640,'Подписчик')
    g.actor(1550,400,'Посетитель')
    g.actor(1550,850,'Модератор')
    for y in (200,430,680,1030):g.line([(175,643),(300,643),(300,y),(385,y)])
    g.uc(600,200,'UC-05. Публикация\nобъявления')
    g.uc(1100,200,'UC-06. Проверка\nобъявления')
    g.uc(600,430,'UC-07. Поиск вещи')
    g.uc(1100,430,'UC-08. Инструкция\nдля метро')
    g.uc(600,680,'UC-09. Заявка\nна возврат')
    g.uc(600,1030,'UC-10. Подтверждение\nпередачи')
    g.uc(1100,850,'UC-11. Жалоба\nи разбор спора')
    g.line([(815,200),(885,200)],dashed=True,arrow=True)
    g.text(850,155,'«include»',20)
    g.line([(1315,430),(1440,430),(1440,403),(1515,403)])
    g.line([(1100,372),(1100,315),(600,315),(600,258)],dashed=True,arrow=True)
    g.text(840,292,'«extend» [место — метро]',20)
    g.line([(1100,908),(1100,1150),(300,1150),(300,643),(175,643)])
    g.line([(1315,200),(1410,200),(1410,853),(1515,853)])
    g.line([(1315,850),(1515,850)])
    g.line([(885,820),(815,710)],dashed=True,arrow=True)
    g.text(995,735,'«extend»\n[спор]',20)
    g.line([(885,880),(815,1000)],dashed=True,arrow=True)
    g.text(980,965,'«extend»\n[спор]',20)
    g.save()
    uml_file('02_return','''actor "Подписчик" as Subscriber
actor "Посетитель" as Visitor
actor "Модератор" as Moderator
rectangle "ИС Потеряшки" {
usecase "UC-05. Публикация объявления" as U5
usecase "UC-06. Проверка объявления" as U6
usecase "UC-07. Поиск вещи" as U7
usecase "UC-08. Инструкция для метро" as U8
usecase "UC-09. Заявка на возврат" as U9
usecase "UC-10. Подтверждение передачи" as U10
usecase "UC-11. Жалоба и разбор спора" as U11
}
Subscriber -- U5
Subscriber -- U7
Subscriber -- U9
Subscriber -- U10
Subscriber -- U11
Visitor -- U8
Moderator -- U6
Moderator -- U11
U5 ..> U6 : <<include>>
U8 ..> U5 : <<extend>>\n[место — метро]
U11 ..> U9 : <<extend>>\n[спор]
U11 ..> U10 : <<extend>>\n[спор]''')

    g=Drawing('03_auction',1720,1400,'ИС «Потеряшки»: аукцион и управление')
    g.box(345,85,1350,1345)
    g.text(845,110,'Граница информационной системы',22)
    g.actor(140,510,'Подписчик')
    g.actor(1550,200,'Модератор')
    g.actor(1550,1130,'Администратор')
    for y in (220,510,800):g.line([(175,513),(300,513),(300,y),(385,y)])
    g.uc(600,220,'UC-12. Допуск\nи выставление лота')
    g.uc(1100,220,'Проверить право\nраспоряжения')
    g.uc(600,510,'UC-13. Ставка')
    g.uc(600,800,'UC-14. Завершение\nаукциона')
    g.uc(1100,1090,'UC-15. Справочник,\nтариф и доступ')
    g.uc(1100,1260,'UC-16. Просмотр\nаудита')
    g.line([(815,220),(885,220)],dashed=True,arrow=True)
    g.text(850,170,'«include»',20)
    g.line([(1315,220),(1440,220),(1440,203),(1515,203)])
    g.line([(1515,1133),(1410,1133),(1410,1090),(1315,1090)])
    g.line([(1515,1133),(1410,1133),(1410,1260),(1315,1260)])
    g.text(1090,800,'UC-14 запускается\nвнутренним планировщиком.\nСтороны получают результат.\nПередача описана в UC-10.',25)
    g.save()
    uml_file('03_auction','''actor "Подписчик" as Subscriber
actor "Модератор" as Moderator
actor "Администратор" as Admin
rectangle "ИС Потеряшки" {
usecase "UC-12. Допуск и выставление лота" as U12
usecase "Проверить право распоряжения" as Permission
usecase "UC-13. Ставка" as U13
usecase "UC-14. Завершение аукциона" as U14
usecase "UC-15. Справочник, тариф и доступ" as U15
usecase "UC-16. Просмотр аудита" as U16
}
Subscriber -- U12
Subscriber -- U13
Subscriber -- U14
Moderator -- Permission
Admin -- U15
Admin -- U16
U12 ..> Permission : <<include>>
note right of U14
Запуск внутренним планировщиком.
Стороны получают результат.
Передача описана в UC-10.
end note''')

    g=Drawing('04_architecture',1720,1390,'Компоненты и развёртывание на helios')
    g.box(65,120,405,355,'Браузер\nNext.js • TypeScript\nRedux Toolkit',BLUE)
    g.box(505,100,1340,1180)
    g.text(920,135,'helios — учётная запись студента',27,bold=True)
    g.box(550,190,1300,925)
    g.text(925,230,'Одна JVM: Spring Boot / Spring MVC',27,bold=True)
    g.box(605,285,1245,390,'REST API • Spring MVC\nSpring Security • JWT\nРаздача frontend Next.js',BLUE,size=24)
    g.box(605,455,1245,640,'Прикладные сервисы\nАккаунты • Подписки • Объявления\nВозвраты • Аукционы • Модерация\nСправочник • Уведомления • Аудит\nРежим demo: тестовые адаптеры',BLUE,size=24)
    g.box(605,705,1245,805,'Hibernate / JPA / JDBC\nФункции PL/pgSQL',BLUE)
    g.box(570,985,920,1125,'Статический frontend\nHTML / CSS / JS\nВ составе JAR',BLUE,size=24)
    g.box(1450,985,1660,1205,'PostgreSQL\npg / studs\ns465826\nтаблицы lf_*',BLUE,size=25)
    g.box(1420,175,1675,370,'Платёжный сервис\nРеальная оплата\n(рабочий режим)',BLUE,size=23)
    g.box(1420,465,1675,650,'SMTP-сервис\nДоставка писем\n(рабочий режим)',BLUE,size=23)
    g.box(65,570,405,860,'MinIO\nФотографии и материалы\nЗакрытый бакет\nОтдельный сервис',BLUE,size=24)
    g.box(65,940,405,1180,'Официальный сайт\nметрополитена\nИсточник контактов\nРучное обновление',BLUE,size=24)
    g.line([(405,335),(605,335)],arrow=True)
    g.text(485,275,'SSH-туннель',21)
    g.text(485,310,'HTTP',21)
    g.line([(925,390),(925,455)],arrow=True)
    g.line([(925,640),(925,705)],arrow=True)
    g.line([(605,600),(450,600),(450,720),(405,720)],arrow=True)
    g.text(235,825,'S3 API',20)
    g.line([(1245,755),(1390,755),(1390,1075),(1450,1075)],arrow=True)
    g.text(1390,925,'JDBC',21)
    g.line([(1245,505),(1370,505),(1370,275),(1420,275)],arrow=True)
    g.line([(1245,565),(1420,565)],arrow=True)
    g.line([(405,1050),(485,1050),(485,540),(605,540)],dashed=True,arrow=True)
    g.text(930,1260,'Локальная сборка: Gradle Wrapper и pnpm.\nНа helios — JAR с frontend; MinIO — отдельный сервис.',26)
    g.save()
    uml_file('04_architecture','''node "Рабочее место" { component "Браузер\nNext.js / TypeScript / Redux" as Browser }
node "helios: учётная запись студента" {
node "Одна JVM Spring Boot" {
component "REST API / Spring MVC / Spring Security / JWT" as Web
component "Прикладные сервисы\nРежим demo: тестовые адаптеры" as Services
component "Hibernate / JPA / JDBC / функции PL/pgSQL" as Repo
}
artifact "Статический frontend\nHTML / CSS / JS в JAR" as Frontend
}
database "PostgreSQL\npg / studs / s465826\nтаблицы lf_*" as DB
node "Отдельный сервис" { database "MinIO\nзакрытый бакет" as Files }
component "Платёжный сервис\nрабочий режим" as Pay
component "SMTP-сервис\nрабочий режим" as Mail
component "Официальный сайт метро\nисточник справочника" as Metro
Browser --> Web : HTTP через SSH-туннель
Web --> Services
Services --> Repo
Repo --> DB : JDBC
Services --> Files : S3 API
Services --> Pay
Services --> Mail
Metro ..> Services : ручное обновление администратором''')


def blocks(text):
    lines=text.splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('```'):
            code=[];i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]);i+=1
            yield 'code','\n'.join(code)
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',v or '-') for v in cells):rows.append(cells)
                i+=1
            yield 'table',rows
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            yield 'heading',(level,line[level:].strip())
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            yield 'image',m.groups()
        elif re.match(r'^Таблица \d+ — ',line):yield 'caption',line
        elif re.match(r'^\d+\. ',line):yield 'number',line
        elif line.startswith('- '):yield 'bullet',line[2:]
        else:
            parts=[line]
            while i+1<len(lines) and lines[i+1].strip() and not re.match(r'^(#|\||!\[|- |\d+\. )',lines[i+1].strip()):
                i+=1;parts.append(lines[i].strip())
            yield 'paragraph',' '.join(parts)
        i+=1


def plain(s):return re.sub(r'\*\*(.*?)\*\*',r'\1',s).replace('<br>','\n')


def word_inline(p,s):
    for k,t in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
        if t:p.add_run(t).bold=bool(k%2)


TITLE_FIELDS=['Университет ИТМО','Факультет программной инженерии и компьютерной техники\nОбразовательная программа\n«Системное и прикладное программное обеспечение»',
              'Курсовая работа\nПо дисциплине «Информационные системы»',
              'Информационная система для поиска\nи возврата потерянных вещей «Потеряшки»',
              'Этап 1',
              'Выполнил студент группы [номер группы]\nЕвграфов Артём Андреевич',
              'Проверил:\n[ФИО преподавателя]', 'Санкт-Петербург 2026']
REPORT_TITLE='ИС «Потеряшки». Отчёт по этапу 1'


def build_docx(items):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin,sec.bottom_margin=Cm(2),Cm(2)
    sec.left_margin,sec.right_margin=Cm(3),Cm(1.5)
    sec.header_distance=sec.footer_distance=Cm(1)
    sec.different_first_page_header_footer=True
    numbering=OxmlElement('w:pgNumType');numbering.set(qn('w:start'),'0');sec._sectPr.append(numbering)
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman';normal.font.size=Pt(14)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Cm(1.25)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ['Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14);st.font.bold=True
        st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                fonts.attrib.pop(qn('w:'+attribute),None)
        st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=True
    doc.styles['Heading 1'].paragraph_format.page_break_before=False
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.paragraph_format.first_line_indent=Cm(0)
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    for i,text in enumerate(TITLE_FIELDS):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1
        p.paragraph_format.space_after=Pt(16)
        if i==2:p.paragraph_format.space_before=Cm(4)
        if i==5:p.paragraph_format.space_before=Cm(2)
        if i in [5,6]:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=p.add_run(text);r.bold=i in [2,3];r.font.size=Pt(18 if i==0 else 16)
        if i==7:p.paragraph_format.space_before=Cm(3.4)
    doc.add_page_break()
    p=doc.add_paragraph('Содержание');p.runs[0].bold=True;p.paragraph_format.first_line_indent=Cm(0)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');p._p.append(toc)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
    for kind,val in items:
        if kind=='heading':
            level,title=val
            if level==1:continue
            p=doc.add_heading(title,level=level-1)
            if title=='Задание':p.paragraph_format.page_break_before=True
        elif kind=='code':
            for line in val.splitlines():
                p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
                r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9)
        elif kind=='caption':
            p=doc.add_paragraph(val);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
            p.paragraph_format.space_after=Pt(8)
            for r in p.runs:r.font.size=Pt(12)
        elif kind in ['paragraph','bullet','number']:
            p=doc.add_paragraph()
            if kind in ['bullet','number']:
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.left_indent=Cm(.4)
            word_inline(p,('• ' if kind=='bullet' else '')+val)
        elif kind=='image':
            caption,path=val
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(REPORT_DIR/path),width=Cm(16.4))
            p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
            for r in p.runs:r.font.size=Pt(12)
        elif kind=='table':
            is_case=val[0][0]=='Прецедент'
            t=doc.add_table(rows=1,cols=len(val[0]));t.style='Table Grid'
            for j,s in enumerate(val[0]):t.rows[0].cells[j].text=plain(s)
            trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
            for row in val[1:]:
                cells=t.add_row().cells
                for j,s in enumerate(row):cells[j].text=plain(s)
            # Match readable table widths in the PDF version.
            widths=table_widths(val,16.5)
            t.autofit=False
            for j,w in enumerate(widths):t.columns[j].width=Cm(w)
            for ri,row in enumerate(t.rows):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for j,cell in enumerate(row.cells):
                    cell.width=Cm(widths[j])
                    for p in cell.paragraphs:
                        p.paragraph_format.first_line_indent=Cm(0)
                        p.paragraph_format.line_spacing=1
                        p.paragraph_format.space_after=Pt(4)
                        p.paragraph_format.space_before=Pt(3)
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.keep_with_next=ri==0 or (is_case and ri==len(val)-1)
                        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=ri==0
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF3F8');cell._tc.get_or_add_tcPr().append(shade)
            if not is_case:doc.add_paragraph().paragraph_format.space_after=Pt(4)
    doc.core_properties.title=REPORT_TITLE
    doc.core_properties.subject='Анализ предметной области, требования, прецеденты, архитектура'
    doc.core_properties.author='Евграфов Артём Андреевич'
    doc.core_properties.comments=''
    doc.core_properties.last_modified_by=''
    doc.core_properties.created=datetime.now(timezone.utc)
    doc.core_properties.modified=datetime.now(timezone.utc)
    path=REPORT_DIR/'report.docx'
    doc.save(path)
    # Remove irrelevant template statistics and application defaults.
    with ZipFile(path) as source:
        parts={name:source.read(name) for name in source.namelist()}
    thumbnails=[name for name in parts if name.startswith('docProps/thumbnail.')]
    for name in thumbnails:del parts[name]
    for name in ['_rels/.rels','[Content_Types].xml']:
        xml=etree.fromstring(parts[name])
        for child in list(xml):
            if 'thumbnail' in child.attrib.get('Type','') or child.attrib.get('PartName','').lstrip('/') in thumbnails:
                xml.remove(child)
        parts[name]=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    xml=etree.fromstring(parts['docProps/app.xml'])
    ns={'ep':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}
    for tag in ['Application','AppVersion','Pages','Words','Characters','Lines','Paragraphs','CharactersWithSpaces','HeadingPairs','TitlesOfParts']:
        element=xml.find('ep:'+tag,ns)
        if element is not None:xml.remove(element)
    parts['docProps/app.xml']=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
    with ZipFile(path,'w',ZIP_DEFLATED) as output:
        for name,data in parts.items():output.writestr(name,data)


def table_widths(rows,total):
    cols=len(rows[0])
    if cols==4:
        if rows[0][0]=='ID':parts=[.10,.31,.43,.16]
        else:parts=[1/4]*4
    elif cols==3:
        parts=[.12,.36,.52] if rows[0][0]=='ID' else [.28,.32,.40]
    else:parts=[.30,.70]
    return [v*total for v in parts]


def pdf_inline(s):
    links=[]
    def named_link(match):
        marker=f'__REPORT_LINK_{len(links)}__'
        links.append(f'<link href="{html.escape(match.group(2),quote=True)}" color="#234e76">{html.escape(match.group(1))}</link>')
        return marker
    s=re.sub(r'\[([^\]]+)\]\((https?://[^\s)]+)\)',named_link,s)
    s=html.escape(s)
    s=s.replace('&lt;br&gt;','<br/>')
    s=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',s)
    s=re.sub(r'https://[^\s<]+',lambda m:f'<link href="{m.group(0)}" color="#234e76">{m.group(0)}</link>',s)
    for i,link in enumerate(links):s=s.replace(f'__REPORT_LINK_{i}__',link)
    return s


def build_pdf(items, toc_pages=None):
    for name,file in [('TimesRU','times.ttf'),('TimesRU-Bold','timesbd.ttf'),('TimesRU-Italic','timesi.ttf'),('TimesRU-BoldItalic','timesbi.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(FONTS/file)))
    pdfmetrics.registerFontFamily('TimesRU',normal='TimesRU',bold='TimesRU-Bold',italic='TimesRU-Italic',boldItalic='TimesRU-BoldItalic')
    pdfmetrics.registerFont(TTFont('MonoRU',str(FONTS/'consola.ttf')))
    base=ParagraphStyle('body',fontName='TimesRU',fontSize=14,leading=21,alignment=TA_JUSTIFY,firstLineIndent=12.5*mm,spaceAfter=3)
    h1=ParagraphStyle('h1',parent=base,fontName='TimesRU-Bold',firstLineIndent=0,alignment=TA_LEFT,spaceAfter=12,keepWithNext=True)
    h2=ParagraphStyle('h2',parent=h1,spaceBefore=12,spaceAfter=6)
    liststyle=ParagraphStyle('list',parent=base,firstLineIndent=0,leftIndent=4*mm)
    ts=ParagraphStyle('table',parent=base,fontSize=11,leading=13,firstLineIndent=0,alignment=TA_LEFT,splitLongWords=True)
    th=ParagraphStyle('tablehead',parent=ts,fontName='TimesRU-Bold')
    caption=ParagraphStyle('caption',parent=base,fontSize=12,leading=15,firstLineIndent=0,alignment=TA_CENTER,spaceAfter=12)
    title=ParagraphStyle('title',parent=base,firstLineIndent=0,alignment=TA_CENTER,fontSize=14,leading=21,spaceAfter=16)
    story=[]
    for i,t in enumerate(TITLE_FIELDS):
        if i==2:story.append(Spacer(1,40*mm))
        if i==5:story.append(Spacer(1,20*mm))
        if i==7:story.append(Spacer(1,34*mm))
        coverstyle=ParagraphStyle('cover'+str(i),parent=title,fontSize=18 if i==0 else 16,leading=22,alignment=TA_RIGHT if i in [5,6] else TA_CENTER)
        story.append(Paragraph(('<b>'+html.escape(t)+'</b>' if i in [2,3] else html.escape(t)).replace('\n','<br/>'),coverstyle))
    story += [PageBreak(),Paragraph('Содержание',h1)]
    for kind,val in items:
        if kind=='heading' and val[0]==2:
            label=pdf_inline(val[1])
            if toc_pages:
                label+=' ................................ '+str(toc_pages[val[1]])
            story.append(Paragraph(label,ParagraphStyle('toc',parent=base,firstLineIndent=0,alignment=TA_LEFT,spaceAfter=5,leading=18)))
    for kind,val in items:
        if kind=='heading':
            level,text=val
            if level==1:continue
            if level==2 and text.startswith('Задание'):story.append(PageBreak())
            p=Paragraph(pdf_inline(text),h1 if level==2 else h2)
            if level==2:p.section_title=text
            story.append(p)
        elif kind in ['paragraph','number','bullet']:
            story.append(Paragraph(pdf_inline(('• ' if kind=='bullet' else '')+val),liststyle if kind in ['number','bullet'] else base))
        elif kind=='code':
            story.append(Preformatted(val,ParagraphStyle('code',fontName='MonoRU',fontSize=9,leading=11,spaceAfter=8),maxLineLength=91))
        elif kind=='caption':story.append(Paragraph(pdf_inline(val),caption))
        elif kind=='image':
            capt,path=val
            im=Image.open(REPORT_DIR/path)
            width=164*mm;height=width*im.height/im.width
            story.append(KeepTogether([PDFImage(str(REPORT_DIR/path),width=width,height=height),Paragraph(pdf_inline(capt),caption)]))
        elif kind=='table':
            is_case=val[0][0]=='Прецедент'
            data=[[Paragraph(pdf_inline(s),th if i==0 else ts) for s in row] for i,row in enumerate(val)]
            table=Table(data,colWidths=table_widths(val,165*mm),repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.5,colors.HexColor('#6c7a85')),('BACKGROUND',(0,0),(-1,0),colors.HexColor(BLUE)),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
            if is_case:
                table.keepWithNext=True
                story.append(table)
            else:story += [table,Spacer(1,5*mm)]
    def footer(canvas,doc):
        canvas.setCreator('')
        canvas.setSubject('Анализ предметной области и проектирование информационной системы')
        canvas._doc.info.producer=''
        if doc.page>1:
            canvas.saveState();canvas.setFont('TimesRU',12);canvas.drawCentredString(A4[0]/2,12*mm,str(doc.page-1));canvas.restoreState()
    class ReportTemplate(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            if hasattr(flowable, 'section_title'):
                self.section_pages[flowable.section_title]=self.page-1
                key='section_'+str(len(self.section_pages))
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.section_title,key,level=0)
    doc=ReportTemplate(str(REPORT_DIR/'report.pdf'),pagesize=A4,leftMargin=30*mm,rightMargin=15*mm,topMargin=20*mm,bottomMargin=20*mm,title=REPORT_TITLE,author='Евграфов Артём Андреевич',allowSplitting=True)
    doc.section_pages={}
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    if toc_pages is None:build_pdf(items,doc.section_pages)


if __name__=='__main__':
    make_diagrams()
    items=list(blocks((REPORT_DIR/'report.md').read_text(encoding='utf-8')))
    start=next(i for i,(kind,val) in enumerate(items) if kind=='heading' and val[0]==2)
    items=items[start:]
    build_docx(items)
    build_pdf(items)
    print('Created DOCX, PDF, 4 PNG, 4 SVG and 4 PlantUML sources.')
